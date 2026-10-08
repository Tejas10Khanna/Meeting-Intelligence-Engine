"""Stage 3: minutes, decisions and action items with LLM-B (+ grounding and verification)."""
import re

from pydantic import BaseModel, Field
from rapidfuzz import fuzz

from .llm import chat_json
from .prompts import DOC_BASIC_SYS, DOC_GROUNDED_SYS, SUMMARY_SYS, VERIFY_SYS

U = "unspecified"


class Topic(BaseModel):
    title: str
    points: list[str] = Field(default_factory=list)


class Decision(BaseModel):
    text: str
    evidence: str = ""
    timestamp: str = ""


class ActionItem(BaseModel):
    task: str
    owner: str = U
    deadline: str = U
    evidence: str = ""
    timestamp: str = ""


class MeetingRecord(BaseModel):
    title: str = "Meeting"
    summary: str = ""
    minutes: list[Topic] = Field(default_factory=list)
    decisions: list[Decision] = Field(default_factory=list)
    open_proposals: list[str] = Field(default_factory=list)
    action_items: list[ActionItem] = Field(default_factory=list)


def fmt_time(t):
    t = int(t)
    return f"{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}"


def transcript_lines(units, refined):
    return [f"[{fmt_time(u.start)}] " + (f"{u.speaker}: " if u.speaker else "") + refined.get(u.id, u.text)
            for u in units]


def chunk_lines(lines, max_words=1800):
    chunks, cur, n = [], [], 0
    for ln in lines:
        w = len(ln.split())
        if cur and n + w > max_words:
            chunks.append(cur)
            cur, n = [], 0
        cur.append(ln)
        n += w
    if cur:
        chunks.append(cur)
    return chunks


def _n(s):
    return re.sub(r"[^\w\s]", "", (s or "").lower())


def _is_u(s):
    return (s or "").strip().lower() in ("", U, "none", "n/a", "unknown", "not specified", "unassigned")


def ground(items, chunk_text, kind, thr=80):
    """Drop items whose quote is not in the transcript; reset owners/deadlines the text does not support."""
    kept, dropped = [], []
    ct = _n(chunk_text)
    for it in items:
        ev = _n(it.get("evidence", ""))
        if not ev or fuzz.partial_ratio(ev, ct) < thr:
            dropped.append({**it, "drop_reason": "evidence_not_in_transcript"})
            continue
        if kind == "action":
            owner, basis = it.get("owner", U), it.get("owner_basis", U)
            if _is_u(owner):
                it["owner"] = U
            elif basis == "named" and _n(owner) not in ct:
                it["owner"] = U
            elif basis == "self" and not re.match(r"speaker \d+", owner.lower()):
                it["owner"] = U
            elif basis not in ("named", "self"):
                it["owner"] = U
            dl = it.get("deadline", U)
            if _is_u(dl) or fuzz.partial_ratio(_n(dl), ct) < 80:
                it["deadline"] = U
        kept.append(it)
    return kept, dropped


def _extract(chunk_text, llm, grounded, context, k, n, temperature=0.0):
    if not grounded:
        return chat_json(llm, DOC_BASIC_SYS, f"Transcript:\n{chunk_text}", max_new_tokens=2500)
    user = f"Meeting context: {context or 'none'}\nThis is part {k} of {n} of the transcript.\n\nTranscript:\n{chunk_text}"
    return chat_json(llm, DOC_GROUNDED_SYS, user, max_new_tokens=3000, temperature=temperature)


def _dedupe(texts, thr=85):
    out = []
    for t in texts:
        if not any(fuzz.token_set_ratio(_n(t), _n(o)) >= thr for o in out):
            out.append(t)
    return out


def _consensus(runs, key, thr=70):
    """Keep items found (fuzzily) in every run."""
    first, rest = runs[0], runs[1:]
    return [it for it in first
            if all(any(fuzz.token_set_ratio(_n(it[key]), _n(o[key])) >= thr for o in r) for r in rest)]


def _verify(llm, decisions, actions):
    items = [{"id": f"d{i}", "claim": d["text"], "evidence": d.get("evidence", "")} for i, d in enumerate(decisions)]
    items += [{"id": f"a{i}", "claim": a["task"], "owner": a.get("owner"), "deadline": a.get("deadline"),
               "evidence": a.get("evidence", "")} for i, a in enumerate(actions)]
    if not items:
        return decisions, actions
    try:
        out = chat_json(llm, VERIFY_SYS, str(items).replace("'", '"'), max_new_tokens=1500)
        bad = {v["id"] for v in out.get("verdicts", []) if v.get("supported") is False}
    except Exception:
        return decisions, actions
    return ([d for i, d in enumerate(decisions) if f"d{i}" not in bad],
            [a for i, a in enumerate(actions) if f"a{i}" not in bad])


def build_record(units, refined, llm, cfg, context=""):
    """cfg: grounded, verify, consensus_runs, chunk_words. Returns (MeetingRecord, report)."""
    grounded = cfg.get("grounded", False)
    lines = transcript_lines(units, refined)
    chunks = chunk_lines(lines, cfg.get("chunk_words", 1800))
    topics, decisions, actions, proposals, dropped = [], [], [], [], []
    summaries = []
    for k, ch in enumerate(chunks, 1):
        text = "\n".join(ch)
        runs = max(1, cfg.get("consensus_runs", 1)) if grounded else 1
        parsed = []
        for r in range(runs):
            try:
                parsed.append(_extract(text, llm, grounded, context, k, len(chunks), 0.0 if r == 0 else 0.3))
            except Exception as e:
                dropped.append({"chunk": k, "error": str(e)[:200]})
        if not parsed:
            continue
        p = parsed[0]
        d_items, a_items = p.get("decisions", []), p.get("action_items", [])
        if not grounded:                                   # normalise the loose baseline format
            d_items = [{"text": d if isinstance(d, str) else d.get("text", ""), "evidence": ""} for d in d_items]
            a_items = [{"task": a.get("task", ""), "owner": a.get("owner", U), "deadline": a.get("deadline", U),
                        "evidence": ""} for a in a_items if isinstance(a, dict)]
        else:
            if len(parsed) > 1:
                d_items = _consensus([q.get("decisions", []) for q in parsed], "text")
                a_items = _consensus([q.get("action_items", []) for q in parsed], "task")
            d_items, d_drop = ground(d_items, text, "decision")
            a_items, a_drop = ground(a_items, text, "action")
            dropped += d_drop + a_drop
            if cfg.get("verify"):
                d_items, a_items = _verify(llm, d_items, a_items)
        topics += [Topic(title=t.get("title", ""), points=t.get("points", [])) for t in p.get("minutes", [])
                   if isinstance(t, dict)]
        decisions += d_items
        actions += a_items
        proposals += [x for x in p.get("open_proposals", []) if isinstance(x, str)]
        if p.get("summary"):
            summaries.append(str(p["summary"]))
    title, summary = "Meeting", " ".join(summaries)
    if len(summaries) > 1:
        try:
            s = chat_json(llm, SUMMARY_SYS, "\n".join(f"Part {i + 1}: {t}" for i, t in enumerate(summaries)))
            title, summary = s.get("title", title), s.get("summary", summary)
        except Exception:
            pass
    elif summaries and cfg.get("title_from_summary", True):
        title = " ".join(summaries[0].split()[:8]) + "…"
    rec = MeetingRecord(
        title=title, summary=summary, minutes=topics,
        decisions=[Decision(text=d["text"], evidence=d.get("evidence", "")) for d in _dedupe_dicts(decisions, "text")],
        open_proposals=_dedupe(proposals),
        action_items=[ActionItem(task=a["task"], owner=a.get("owner", U) or U, deadline=a.get("deadline", U) or U,
                                 evidence=a.get("evidence", "")) for a in _dedupe_dicts(actions, "task")])
    return rec, {"dropped_items": dropped, "chunks": len(chunks)}


def _dedupe_dicts(items, key, thr=85):
    out = []
    for it in items:
        if it.get(key) and not any(fuzz.token_set_ratio(_n(it[key]), _n(o[key])) >= thr for o in out):
            out.append(it)
    return out


def to_markdown(r: MeetingRecord):
    L = [f"# {r.title}", "", "## Summary", r.summary or "_No summary._", "", "## Minutes"]
    for t in r.minutes:
        L += [f"### {t.title}"] + [f"- {p}" for p in t.points]
    L += ["", "## Key decisions"] + ([f"{i}. {d.text}" for i, d in enumerate(r.decisions, 1)] or ["_None recorded._"])
    L += ["", "## Open proposals (not agreed)"] + ([f"- {p}" for p in r.open_proposals] or ["_None._"])
    L += ["", "## Action items"]
    L += [f"{i}. **{a.task}** — Owner: {a.owner}; Deadline: {a.deadline}" for i, a in enumerate(r.action_items, 1)] \
        or ["_None recorded._"]
    return "\n".join(L)


def to_docx(r: MeetingRecord, path):
    import docx
    d = docx.Document()
    d.add_heading(r.title, 0)
    d.add_heading("Summary", 1)
    d.add_paragraph(r.summary)
    d.add_heading("Minutes", 1)
    for t in r.minutes:
        d.add_heading(t.title, 2)
        for p in t.points:
            d.add_paragraph(p, style="List Bullet")
    d.add_heading("Key decisions", 1)
    for x in r.decisions:
        d.add_paragraph(x.text, style="List Number")
    d.add_heading("Open proposals (not agreed)", 1)
    for x in r.open_proposals:
        d.add_paragraph(x, style="List Bullet")
    d.add_heading("Action items", 1)
    tb = d.add_table(rows=1, cols=3)
    for c, h in zip(tb.rows[0].cells, ("Task", "Owner", "Deadline")):
        c.text = h
    for a in r.action_items:
        row = tb.add_row().cells
        row[0].text, row[1].text, row[2].text = a.task, a.owner, a.deadline
    d.save(path)
